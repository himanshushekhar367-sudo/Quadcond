#!/usr/bin/env python3
"""Measure what a candidate collection actually adds, before any of it is used.

Written for G4All (bioRxiv 10.64898/2026.08.05.743020, Zenodo 10.5281/zenodo.21216943),
which is a curated compilation of the published literature. That is exactly what
makes it dangerous to ingest quickly: the atlas already holds the G4STAB
stability collection, the G4ShapePredictor topology collection and the
iM-Seeker transitional-pH collection, and a literature compilation is very
likely to contain the same measurements. Two consequences, in order of severity.

1. **Benchmark contamination.** The published-tool benchmark evaluates QuadCond
   and every competitor on those collections. If sequences from them re-enter
   training through a differently-named file, every headline figure in the
   manuscript silently becomes partly in-sample, and nothing in the pipeline
   would report it.
2. **Fold leakage.** The same measurement arriving twice under two source DOIs
   lands in two grouped folds, which inflates cross-validated performance in a
   way that looks like a better model.

So the first question is not "how do we ingest this" but "how much of it is
new". This answers that and writes a quarantine table of the rows that are, with
nothing merged and nothing retrained.

    python g4all_overlap.py /path/to/G4All.xlsx

Note the file distributed as `G4All.csv` is reportedly an Excel workbook; the
reader sniffs the real format rather than trusting the extension.
"""
from __future__ import annotations

import argparse
import csv
import hashlib
import json
import re
import sqlite3
import sys
import zipfile
from collections import Counter
from pathlib import Path

ROOT = Path(__file__).resolve().parent
REPO = ROOT.parent.parent
SEQ_OK = re.compile(r"^[ACGT]+$")


def norm(seq: str) -> str | None:
    """One spelling per molecule: upper case, RNA to DNA, no whitespace.

    Comparing raw strings would count 'gggttagggttagggttaggg' and its uppercase
    twin as two different sequences and report an overlap of zero, which is the
    comfortable answer and the wrong one.
    """
    if not seq:
        return None
    s = re.sub(r"\s+", "", str(seq)).upper().replace("U", "T")
    return s if SEQ_OK.match(s) else None


def seq_hash(s: str) -> str:
    return hashlib.sha256(s.encode()).hexdigest()


def sniff_and_read(path: Path) -> list[dict]:
    """Read the candidate file whatever it actually is."""
    head = path.open("rb").read(8)
    if head[:2] == b"PK":                       # Excel Open XML, whatever it is named
        try:
            from openpyxl import load_workbook
        except ImportError:
            raise SystemExit("this file is an Excel workbook; pip install openpyxl")
        wb = load_workbook(path, read_only=True, data_only=True)
        ws = wb[wb.sheetnames[0]]
        rows = ws.iter_rows(values_only=True)
        header = [str(c).strip() if c is not None else "" for c in next(rows)]
        return [dict(zip(header, r)) for r in rows]
    if head[:1] == b"\x1f":
        raise SystemExit("gzipped input; decompress it first")
    with path.open(encoding="utf-8-sig", newline="") as fh:
        sample = fh.read(8192); fh.seek(0)
        try:
            dialect = csv.Sniffer().sniff(sample, delimiters=",\t;")
        except csv.Error:
            dialect = csv.excel
        return list(csv.DictReader(fh, dialect=dialect))


def find_column(rows: list[dict], *candidates: str) -> str | None:
    if not rows:
        return None
    keys = {k.lower().strip(): k for k in rows[0] if k}
    for c in candidates:
        if c in keys:
            return keys[c]
    for k_lower, k in keys.items():
        if any(c in k_lower for c in candidates):
            return k
    return None


def atlas_index(db: Path):
    """seq_hash -> what the atlas already knows about that sequence."""
    con = sqlite3.connect(db)
    idx: dict[str, dict] = {}
    dois: Counter = Counter()
    q = ("select sequence, tm, ph_t, topology, folded, source_doi, evidence_tier "
         "from records")
    for seq, tm, pht, topo, folded, doi, tier in con.execute(q):
        n = norm(seq)
        if not n:
            continue
        e = idx.setdefault(seq_hash(n), {"tm": 0, "pht": 0, "topology": 0, "folded": 0,
                                         "experimental": 0, "dois": set()})
        e["tm"] += tm is not None
        e["pht"] += pht is not None
        e["topology"] += topo is not None
        e["folded"] += folded is not None
        e["experimental"] += (tier == "experimental")
        if doi:
            e["dois"].add(str(doi).lower().strip())
            dois[str(doi).lower().strip()] += 1
    con.close()
    return idx, dois


def benchmark_sets() -> dict[str, set[str]]:
    """Every sequence the published-tool benchmark evaluates on.

    Overlap with these is the contamination that matters: these are the sets the
    manuscript's headline numbers are computed over.
    """
    out: dict[str, set[str]] = {}
    res = REPO / "benchmarks" / "published_tools" / "results"
    tool = res / "tool_scores_all.csv"
    if tool.exists():
        s = set()
        for r in csv.DictReader(tool.open(encoding="utf-8")):
            n = norm(r.get("sequence", ""))
            if n:
                s.add(seq_hash(n))
        out["benchmark evaluation sequences (tool_scores_all)"] = s
    pairs = res / "pairs_g4_tm.csv"
    if pairs.exists():
        s = set()
        for r in csv.DictReader(pairs.open(encoding="utf-8")):
            for col in ("ref", "alt"):
                n = norm(r.get(col, ""))
                if n:
                    s.add(seq_hash(n))
        out["dTm pair set (pairs_g4_tm)"] = s
    return out


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("candidate", help="the downloaded collection (xlsx or csv)")
    ap.add_argument("--atlas", default=None,
                    help="atlas to compare against; defaults to the full atlas.db when present")
    ap.add_argument("--name", default="G4All")
    ap.add_argument("--out", default=str(ROOT / "results"))
    a = ap.parse_args()

    path = Path(a.candidate)
    if not path.exists():
        raise SystemExit(f"not found: {path}")
    rows = sniff_and_read(path)
    if not rows:
        raise SystemExit("no rows parsed; check the file")

    seq_col = find_column(rows, "sequence", "seq")
    if not seq_col:
        raise SystemExit(f"no sequence column found. Columns: {list(rows[0])[:20]}")
    doi_col = find_column(rows, "doi", "reference", "source")
    tm_col = find_column(rows, "tm", "melting")
    topo_col = find_column(rows, "topology", "conformation")
    call_col = find_column(rows, "conclusion", "folding", "g4 status", "class")

    # The full atlas, not the trimmed core. atlas_core.db holds a subset, so
    # auditing against it reports sequences as new that the full atlas already
    # has -- an overlap audit that understates overlap is worse than none.
    atlas_path = Path(a.atlas) if a.atlas else None
    if atlas_path is None:
        full, core = REPO / "data" / "atlas.db", REPO / "data" / "atlas_core.db"
        atlas_path = full if full.exists() else core
        if atlas_path == core:
            print("WARNING: comparing against atlas_core.db, which is a subset of the full "
                  "atlas. Overlap reported below is a LOWER BOUND. Fetch data/atlas.db and "
                  "rerun before acting on these numbers.\n", file=sys.stderr)
    if not atlas_path.exists():
        raise SystemExit(f"no atlas at {atlas_path}")
    atlas, atlas_dois = atlas_index(atlas_path)
    bench = benchmark_sets()

    seen: dict[str, dict] = {}
    unparsable = 0
    for r in rows:
        n = norm(r.get(seq_col, ""))
        if not n:
            unparsable += 1
            continue
        h = seq_hash(n)
        e = seen.setdefault(h, {"sequence": n, "n_rows": 0, "dois": set(), "row": r})
        e["n_rows"] += 1
        if doi_col and r.get(doi_col):
            e["dois"].add(str(r[doi_col]).lower().strip())

    in_atlas = {h for h in seen if h in atlas}
    novel = [h for h in seen if h not in atlas]
    report = {
        "candidate": str(path),
        "atlas_compared_against": atlas_path.name,
        "atlas_is_full": atlas_path.name == "atlas.db",
        "name": a.name,
        "rows_read": len(rows),
        "rows_with_an_unparsable_sequence": unparsable,
        "distinct_sequences": len(seen),
        "already_in_atlas": len(in_atlas),
        "new_to_atlas": len(novel),
        "fraction_new": round(len(novel) / len(seen), 4) if seen else None,
        "columns_used": {"sequence": seq_col, "doi": doi_col, "tm": tm_col,
                         "topology": topo_col, "call": call_col},
    }

    # What the overlapping sequences already carry. A sequence the atlas holds
    # with a melting temperature is a duplicate measurement; one it holds with
    # only a folding call may still add a Tm.
    carries = Counter()
    for h in in_atlas:
        e = atlas[h]
        for k in ("tm", "pht", "topology", "folded", "experimental"):
            if e[k]:
                carries[f"atlas already has {k}"] += 1
    report["overlap_detail"] = dict(carries)

    # The contamination check.
    contamination = {}
    for label, s in bench.items():
        hit = len(s & set(seen))
        contamination[label] = {"benchmark_sequences": len(s), "also_in_candidate": hit,
                                "fraction_of_benchmark_covered": round(hit / len(s), 4) if s else None}
    report["benchmark_overlap"] = contamination
    report["verdict"] = verdict(report)

    outdir = Path(a.out); outdir.mkdir(parents=True, exist_ok=True)
    stem = a.name.lower().replace(" ", "_")
    (outdir / f"{stem}_overlap.json").write_text(json.dumps(report, indent=2), encoding="utf-8")

    # The quarantine table: genuinely new sequences, nothing merged.
    with (outdir / f"{stem}_quarantine.tsv").open("w", newline="", encoding="utf-8") as fh:
        w = csv.writer(fh, delimiter="\t")
        w.writerow(["sequence", "n_rows_in_candidate", "source_dois",
                    tm_col or "tm", topo_col or "topology", call_col or "call"])
        for h in novel:
            e = seen[h]; r = e["row"]
            w.writerow([e["sequence"], e["n_rows"], ";".join(sorted(e["dois"])),
                        r.get(tm_col, ""), r.get(topo_col, ""), r.get(call_col, "")])

    print(json.dumps(report, indent=2))
    print(f"\nwritten: {outdir}/{stem}_overlap.json and {stem}_quarantine.tsv")
    print("Nothing was merged and nothing was retrained. Read the verdict first.")


def verdict(rep: dict) -> list[str]:
    out = []
    worst = max((v["fraction_of_benchmark_covered"] or 0)
                for v in rep["benchmark_overlap"].values()) if rep["benchmark_overlap"] else 0
    if worst > 0.05:
        out.append(
            f"CONTAMINATION RISK: this collection contains {worst:.1%} of a benchmark evaluation "
            "set. Training on it as-is makes the published comparison partly in-sample. Either "
            "exclude every benchmark sequence before training, or rerun the benchmark on sets "
            "that exclude them and report both.")
    else:
        out.append("No material overlap with the benchmark evaluation sets.")
    if rep["fraction_new"] is not None and rep["fraction_new"] < 0.5:
        out.append(
            f"Only {rep['fraction_new']:.1%} of distinct sequences are new to the atlas. Describe "
            "the addition by that number, not by the collection's headline row count.")
    out.append("Duplicate measurements arriving under a second source DOI split across grouped "
               "folds and inflate cross-validated performance. Keep primary-paper provenance on "
               "every ingested row and group on it.")
    return out


if __name__ == "__main__":
    main()
