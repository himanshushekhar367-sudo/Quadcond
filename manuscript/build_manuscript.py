"""Reproduce the manuscript, figures and tables from frozen metadata; no fitting.

Dependencies: python-docx, matplotlib. Run from any directory.
"""
from pathlib import Path
import csv
from decimal import Decimal, ROUND_HALF_UP
import hashlib
import html
import json
import re

import matplotlib
matplotlib.use('Agg')
import matplotlib.pyplot as plt
from matplotlib.patches import FancyBboxPatch
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.text import WD_BREAK

ROOT = Path(__file__).resolve().parent
ART = ROOT / 'inputs' if (ROOT / 'inputs').exists() else ROOT.parent / 'quadcond' / 'artifacts'
MAIN = json.loads((ART / 'quadcond_model.json').read_text(encoding='utf-8'))
ABL = json.loads((ART / 'quadcond_model_seqonly.json').read_text(encoding='utf-8'))
REFS = json.loads((ROOT / 'verified_references.json').read_text(encoding='utf-8'))
HEADS = ['g4_tm', 'im_pht', 'im_pht_condition', 'im_tm_condition']
LABELS = ['G4 melting temperature', 'i-motif transitional pH',
          'i-motif transitional pH, condition model', 'i-motif melting temperature, condition model']
SHORT = ['G4 melting temperature', 'i-motif transitional pH',
         'i-motif pH, condition model', 'i-motif Tm, condition model']
plt.rcParams.update({'font.family': 'DejaVu Sans', 'font.size': 10,
                     'pdf.fonttype': 42, 'ps.fonttype': 42, 'svg.fonttype': 'none'})
FIG = ROOT / 'figures'
FIG.mkdir(exist_ok=True)


def savefig(fig, stem):
    for ext in ['png', 'pdf', 'svg']:
        fig.savefig(FIG / f'{stem}.{ext}', dpi=300, bbox_inches='tight', facecolor='white')
    plt.close(fig)


def workflow():
    fig, ax = plt.subplots(figsize=(10, 6.6))
    ax.set_xlim(0, 10); ax.set_ylim(-.25, 7); ax.axis('off')
    def box(x, y, w, h, title, body, color):
        ax.add_patch(FancyBboxPatch((x, y), w, h, boxstyle='round,pad=0.1,rounding_size=0.1',
                                   linewidth=1.0, edgecolor=color, facecolor='#f7f9fb'))
        ax.text(x+w/2, y+h-0.24, title, ha='center', va='top', color=color, weight='bold', fontsize=11)
        ax.text(x+w/2, y+h/2-0.16, body, ha='center', va='center', fontsize=10, linespacing=1.4)
    def arrow(x1,y1,x2,y2):
        ax.annotate('', xy=(x2,y2), xytext=(x1,y1), arrowprops={'arrowstyle':'-|>', 'color':'#607080','lw':1.3})
    box(2.0,5.65,6,1.1,'Sequence and requested conditions',"Entered strand (+) and reverse complement (−)\nExplicit values and imputed defaults",'#184e72')
    arrow(5,5.52,5,5.12)
    box(2.0,3.9,6,1.1,'Task-specific models and applicability',
        'Supported estimate  |  Extrapolation hidden  |  Refused: no number','#184e72')
    for x in [1.6,5,8.4]:
        arrow(5,3.78,x,3.3)
    box(.2,1.9,2.8,1.25,'Biophysical evidence','Melting temperature\nTransitional pH and class','#226d59')
    box(3.6,1.9,2.8,1.25,'Genomic proxies','Peak-label association\nNo molecular occupancy','#9a6020')
    box(7,1.9,2.8,1.25,'Auxiliary evidence','Predicted or derived labels\nNo new physical measurement','#68527b')
    for x in [1.6,5,8.4]: arrow(x,1.77,5,1.23)
    box(.7,.05,8.6,1.05,'Inspect, compare and export',
        'Evidence cards  •  Condition sweeps  •  Mutation scans  •  Batch records\nSchematic 3D classes; source, model identity and limitations retained','#184e72')
    savefig(fig,'Figure_1_workflow')


def ablation():
    fig, axes = plt.subplots(2,2,figsize=(9.5,6),sharex=True)
    rows=[]
    for i,(ax,key,label) in enumerate(zip(axes.flat,HEADS,SHORT)):
        h=MAIN['heads'][key]; t=h['training_meta']
        values=[ABL['heads'][key]['metrics']['r2'], h['metrics']['r2']]
        ax.barh([1,0],values,height=.48,color=['#99a7b5','#186a7b'])
        ax.set_yticks([1,0],['Sequence only','Sequence + conditions'])
        for y,v in zip([1,0],values): ax.text(v+.015,y,f'{v:.3f}',va='center',fontsize=10)
        ax.set_xlim(0,1.05); ax.set_xticks([0,.25,.5,.75,1]); ax.set_xlabel('Recorded R²')
        ax.set_title(f'{chr(65+i)}  {label}',loc='left',fontsize=11,weight='bold',pad=29)
        scope=f"{t['n_rows']:,} records; {t['grouping']['n_groups']} {t['group_by']} groups"
        ax.text(0,1.1,scope,transform=ax.transAxes,fontsize=9,color='#455565')
        ax.spines[['top','right','left']].set_visible(False)
        ax.grid(axis='x',alpha=.18); ax.set_axisbelow(True)
        rows.append([key,t['n_rows'],t['grouping']['n_groups'],t['group_by'],*values])
    fig.text(.5,.012,'C–D: condition holdout for C9 and Tel21C only. Point estimates; no uncertainty bars.',ha='center',fontsize=9)
    fig.tight_layout(rect=(0,.05,1,1),h_pad=2.5,w_pad=2)
    savefig(fig,'Figure_2_ablation')
    csvout('figure_2_data.csv',['head','records','groups','group_by','sequence_only_r2','sequence_conditions_r2'],rows)


def csvout(name,header,rows):
    with (ROOT/name).open('w',encoding='utf-8-sig',newline='') as f:
        w=csv.writer(f); w.writerow(header); w.writerows(rows)


def _r3(x):
    # Round half up. '%.3f' inherits binary rounding, so 0.4775 prints as 0.477
    # and the prose that quotes 0.478 looks like a different number.
    return str(Decimal(repr(float(x))).quantize(Decimal('0.001'), rounding=ROUND_HALF_UP))


def _cell(x):
    # A literal pipe inside a cell splits that row into more cells than the
    # header has, and the docx writer then indexes past the end of the row.
    # Tool names such as 'G4Hunter (max |25-nt window|)' carry them, so strip
    # them here rather than relying on every caller to remember.
    return str(x).replace('|', '')


def table_md(headers, rows):
    headers=[_cell(h) for h in headers]; rows=[[_cell(c) for c in r] for r in rows]
    return '\n'.join(['| '+' | '.join(headers)+' |', '| '+' | '.join(['---']*len(headers))+' |']+
                     ['| '+' | '.join(map(str,r))+' |' for r in rows])



# ---------------------------------------------------------------- v0.5.1 tables
# Three tables added for the v0.5.1 revision. Each is built from a frozen result
# file under inputs/, so the manuscript cannot drift from the numbers the
# benchmark and the genome-wide run actually produced, and nothing here refits.

BENCH_ROWS = [
    # (task substring, tool exact, metric, display label, note)
    ('G4 melting temperature', 'QuadCond g4_tm (grouped OOF)', 'R2', 'G4 melting temperature (R2, grouped)', 'QuadCond'),
    ('G4 melting temperature', 'G4STAB architecture retrained (grouped 5-fold)', 'R2', 'G4 melting temperature (R2, grouped)', 'G4STAB architecture, retrained on the same folds'),
    ('G4 melting temperature', 'G4STAB released ensemble', 'R2', 'G4 melting temperature (R2, grouped)', 'G4STAB released model (in-sample upper bound)'),
    ('G4 melting temperature', 'QuadCond g4_tm (random 5-fold, for comparison with published G4STAB protocol)', 'R2', 'G4 melting temperature (R2, random split)', 'QuadCond under the published G4STAB protocol'),
    ('Within-sequence buffer response', 'QuadCond g4_tm (grouped OOF)', 'median per-sequence Spearman', 'Buffer response (median per-sequence rho)', 'QuadCond'),
    ('Within-sequence buffer response', 'G4STAB architecture retrained (grouped 5-fold)', 'median per-sequence Spearman', 'Buffer response (median per-sequence rho)', 'G4STAB architecture, retrained'),
    ('Within-sequence buffer response', 'Any buffer-blind tool (G4Hunter, pqsfinder, QGRS, G4Catchall, G4Boost, DeepG4, G4detector, G4mismatch)', 'median per-sequence Spearman', 'Buffer response (median per-sequence rho)', 'any buffer-blind tool (constant by construction)'),
    ('Single-substitution effect', 'QuadCond g4_tm (Δ of grouped OOF predictions)', 'Spearman', 'Single-substitution dTm (rho)', 'QuadCond'),
    ('Single-substitution effect', 'G4STAB architecture retrained, grouped (Δ)', 'Spearman', 'Single-substitution dTm (rho)', 'G4STAB architecture, retrained'),
    ('Single-substitution effect', 'Δpqsfinder', 'Spearman', 'Single-substitution dTm (rho)', 'delta pqsfinder'),
    ('Single-substitution effect', 'ΔDeepG4', 'Spearman', 'Single-substitution dTm (rho)', 'delta DeepG4'),
    ('Single-substitution effect', 'G4SNVHunter G4VarImpact (ΔG4Hunter window score)', 'Spearman', 'Single-substitution dTm (rho)', 'G4SNVHunter G4VarImpact'),
    ('Single-substitution effect', 'QuadCond g4_tm (Δ of grouped OOF predictions)', 'sign accuracy |ΔTm|≥2°C', 'Single-substitution dTm (sign accuracy, |dTm| >= 2 C)', 'QuadCond'),
    ('Single-substitution effect', 'G4SNVHunter G4VarImpact (ΔG4Hunter window score)', 'sign accuracy |ΔTm|≥2°C', 'Single-substitution dTm (sign accuracy, |dTm| >= 2 C)', 'G4SNVHunter G4VarImpact'),
    ('i-motif folding', 'QuadCond im_fold (grouped OOF)', 'AUROC', 'i-motif folding vs dinucleotide shuffles (AUROC)', 'QuadCond'),
    ('i-motif folding', 'iM-Seeker (folding probability)', 'AUROC', 'i-motif folding vs dinucleotide shuffles (AUROC)', 'iM-Seeker'),
    ('i-motif transitional pH', 'QuadCond im_pht on the same 137 sequences', 'Spearman', 'i-motif transitional pH (rho, 137 shared sequences)', 'QuadCond'),
    ('i-motif transitional pH', 'iM-Seeker architecture retrained (grouped folds)', 'Spearman', 'i-motif transitional pH (rho, 137 shared sequences)', 'iM-Seeker architecture, retrained'),
]


def _bench_index():
    rows = {}
    with (ROOT / 'inputs' / 'benchmark_results.csv').open(encoding='utf-8') as fh:
        for r in csv.DictReader(fh):
            rows[(r['task'], r['tool'], r['metric'])] = r
    return rows


def benchmark_table():
    idx = _bench_index()
    out, seen = [], None
    for task_sub, tool, metric, label, note in BENCH_ROWS:
        hit = next((v for (t, to, m), v in idx.items()
                    if task_sub in t and to == tool and m == metric), None)
        assert hit, f'benchmark row not found: {task_sub} | {tool} | {metric}'
        lo, hi = hit['ci_low'], hit['ci_high']
        ci = f"{_r3(lo)}–{_r3(hi)}" if lo and hi else '—'
        out.append([label if label != seen else '', note,
                    _r3(hit['value']), ci, f"{int(hit['n']):,}"])
        seen = label
    return table_md(['Task', 'Model', 'Value', '95% CI', 'n'], out)


SCANNER_TOOLS = ['QuadCond genome-scan (G4-seq)', 'G4Hunter (max |25-nt window|)', 'pqsfinder',
                 'DeepG4', 'canonical G3 regex count', 'G4mismatch (K)',
                 'G4detector (K, random-neg model)', 'G4detector (K, PQ-neg model)']


def scanner_table():
    vals = {}
    with (ROOT / 'inputs' / 'scanner_benchmark.csv').open(encoding='utf-8') as fh:
        for r in csv.DictReader(fh):
            vals[(r['split'], r['negatives'], r['tool'])] = float(r['auroc'])
    rows = []
    for split, label in [('test_heldout_chr', 'Held-out chromosomes 2, 8, 17'), ('test_mouse', 'Mouse (cross-species)')]:
        first = True
        for tool in SCANNER_TOOLS:
            rnd, pq = vals.get((split, 'random', tool)), vals.get((split, 'pq', tool))
            assert rnd is not None and pq is not None, f'scanner row missing: {split} {tool}'
            rows.append([label if first else '', tool, _r3(rnd), _r3(pq), _r3(min(rnd, pq))])
            first = False
    return table_md(['Evaluation set', 'Tool', 'AUROC vs random negatives',
                     'AUROC vs PQS-matched negatives', 'Lower of the two'], rows)


GW_CLASSES = [('motif_lost', 'Motif abolished'), ('destabilising', 'Destabilising'),
              ('moderate', 'Moderate'), ('neutral', 'Neutral'), ('stabilising', 'Stabilising'),
              ('flank', 'Flanking window (reference)'), ('control', 'Distant GC-matched control')]


def genomewide_table():
    summ = json.loads((ROOT / 'inputs' / 'genomewide_summary.json').read_text(encoding='utf-8'))
    rows = []
    for kind in ('G4', 'iM'):
        k = summ['kinds'][kind]
        terms = k['adjusted_rank_model']['terms']
        bal = k['covariate_balance']
        first = True
        for cls, label in GW_CLASSES:
            if cls not in bal:
                continue
            b = bal[cls]
            if cls == 'flank':
                shift = 'reference'
            else:
                t = terms.get(cls)
                shift = (f"{t['delta_pctile'] * 100:+.2f} ({t['ci'][0] * 100:+.2f} to "
                         f"{t['ci'][1] * 100:+.2f}), p = {t['p']:.2g}") if t else '—'
            rows.append([kind if first else '', label, f"{b['n']:,}",
                         f"{b['mean_gc']:.3f}", f"{b['frac_cpg'] * 100:.1f}%", shift])
            first = False
    return table_md(['Motif class', 'Substitution class', 'n', 'Mean GC', 'CpG',
                     'Percentile shift vs flank (95% CI)'], rows)

def references(text):
    order=[]
    def cite(match):
        keys=[x.strip().removeprefix('@') for x in match.group(1).split(';')]
        for k in keys:
            assert k in REFS, f'Unknown reference: {k}'
            if k not in order: order.append(k)
        return '('+', '.join(map(str, sorted({order.index(k)+1 for k in keys})))+')'
    text=re.sub(r'\[(@[a-z0-9_]+(?:;@[a-z0-9_]+)*)\]',cite,text)
    refs=[]; bib=[]
    for n,key in enumerate(order,1):
        if key=='xgb':
            v=REFS[key]['crossref']; doi=v['DOI']
            title='XGBoost: A Scalable Tree Boosting System.'
            line=f'Chen, T. and Guestrin, C. (2016) {title} Proceedings of the 22nd ACM SIGKDD International Conference on Knowledge Discovery and Data Mining, 785–794. https://doi.org/{doi}'
            bib.append(f'@inproceedings{{{key},\n  author = {{Chen, Tianqi and Guestrin, Carlos}},\n  title = {{{title}}},\n  year = {{2016}},\n  doi = {{{doi}}}\n}}')
        else:
            entry=REFS[key]; v=entry['records'][0]; j=v['journalInfo']; doi=v['doi']
            authors=html.unescape(v['authorString']); title=html.unescape(v['title']); journal=j['journal']['title']
            locator=v.get('pageInfo','')
            volume=j.get('volume','')
            # Locators are omitted when the retrieved record does not carry
            # them, rather than rendered as empty commas. An entry that reads
            # "Journal, , ." looks like a formatting bug; an entry with no
            # volume is a reference whose locator still has to be completed,
            # and reference_audit.json is where that is recorded.
            tail=', '.join(x for x in (volume, locator) if x)
            line=f"{authors} ({v['pubYear']}) {title} {journal}" + (f", {tail}" if tail else "")
            url=v.get('url','')
            # A resource announcement has a URL and no DOI. Emitting
            # "https://doi.org/" with nothing after it is a dead link that
            # looks like a real one.
            line+=(f". https://doi.org/{doi}" if doi else (f". {url}" if url else "."))
            if v.get('pmcid'): line+=f"; https://pmc.ncbi.nlm.nih.gov/articles/{v['pmcid']}/"
            elif v.get('pmid'): line+=f"; https://pubmed.ncbi.nlm.nih.gov/{v['pmid']}/"
            # Europe PMC records carry a structured author list; the scite-store
            # records added for the v0.4.9 revision carry only the author
            # string, which is already in the journal's own order.
            names=' and '.join(a.get('fullName','') for a in v.get('authorList',{}).get('author',[]))
            if not names:
                names=' and '.join(a.strip().rstrip('.') for a in authors.rstrip('.').split(',') if a.strip())
            bib.append(f"@article{{{key},\n  author = {{{names}}},\n  title = {{{title}}},\n  journal = {{{journal}}},\n  year = {{{v['pubYear']}}},\n  volume = {{{volume}}},\n  pages = {{{locator}}},\n  doi = {{{doi}}}\n}}")
        refs.append(f'{n}. {line}')
    (ROOT/'references.bib').write_text('\n\n'.join(bib)+'\n',encoding='utf-8')
    provenance={}
    incomplete=[]
    for key in order:
        entry=REFS.get(key,{})
        provenance[key]=entry.get('source','Europe PMC (resultType=core)')
        if entry.get('source') and not entry.get('locator_complete'):
            incomplete.append({'key':key,'doi':entry['records'][0]['doi'],
                               'note':entry.get('note','')})
    (ROOT/'reference_audit.json').write_text(json.dumps({'order':order,'count':len(order),
        'retrieval_date':'2026-09-10',
        'basis':('Publisher/PubMed metadata retrieved through Europe PMC; XGBoost Crossref and '
                 'primary ACM/arXiv record; the v0.4.9 additions from the scite metadata store, '
                 'because Europe PMC was unreachable from the machine that assembled them.'),
        'provenance_by_key':provenance,
        # Named rather than guessed. A volume or page invented to make a
        # reference list look finished is a fabricated locator, and it is
        # exactly the kind of detail nobody re-checks at proof stage.
        'locators_to_complete_before_submission':incomplete,
        'all_citations_resolved':True,'uncited_records':sorted(set(REFS)-set(order))},indent=2),encoding='utf-8')
    return text.replace('{{REFERENCES}}','\n\n'.join(refs))


def hyperlink(p,url):
    link=OxmlElement('w:hyperlink')
    link.set(qn('r:id'),p.part.relate_to(url,'http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink',is_external=True))
    run=OxmlElement('w:r'); pr=OxmlElement('w:rPr'); color=OxmlElement('w:color'); color.set(qn('w:val'),'185D78'); pr.append(color)
    run.append(pr); t=OxmlElement('w:t'); t.text=url; run.append(t); link.append(run); p._p.append(link)


SUP = re.compile(r'<sup>(.*?)</sup>', re.S)


def add_text(p,text):
    """Runs for one Markdown line: real superscripts, live links, plain text.

    Superscript support exists because the title page needs it. Affiliation and
    equal-contribution markers written as ``<sup>1</sup>`` were emitted into the
    document as literal angle brackets -- the author line of the manuscript read
    ``Auroni Deep <sup>1,†</sup>`` in the rendered PDF. A title page is the one
    page a reader is guaranteed to look at.
    """
    def emit(chunk, superscript=False):
        pos=0
        for m in re.finditer(r'https://\S+',chunk):
            if m.start()>pos:
                r=p.add_run(chunk[pos:m.start()]); r.font.superscript=superscript
            url=m.group().rstrip(';.'); hyperlink(p,url)
            trailing=m.group()[len(url):]
            if trailing:
                r=p.add_run(trailing); r.font.superscript=superscript
            pos=m.end()
        if pos<len(chunk):
            r=p.add_run(chunk[pos:]); r.font.superscript=superscript

    cursor=0
    for m in SUP.finditer(text):
        emit(text[cursor:m.start()])
        emit(m.group(1), superscript=True)
        cursor=m.end()
    emit(text[cursor:])


def docx_from_md(md):
    doc=Document(); sec=doc.sections[0]
    sec.page_width=Inches(8.27); sec.page_height=Inches(11.69)
    sec.top_margin=sec.bottom_margin=Inches(.78); sec.left_margin=sec.right_margin=Inches(.82)
    sec.header_distance=sec.footer_distance=Inches(.35)
    normal=doc.styles['Normal']; normal.font.name='Times New Roman'; normal.font.size=Pt(11)
    normal.paragraph_format.line_spacing=1.15; normal.paragraph_format.space_after=Pt(7)
    for name,size in [('Title',23),('Heading 1',15),('Heading 2',12)]:
        s=doc.styles[name]; s.font.name='Arial'; s.font.size=Pt(size); s.font.color.rgb=RGBColor.from_string('183E50')
        s.paragraph_format.space_before=Pt(12); s.paragraph_format.space_after=Pt(7)
    head=sec.header.paragraphs[0]; head.text='QuadCond and AENNA-3D | Author working draft'; head.style='Caption'
    foot=sec.footer.paragraphs[0]; foot.alignment=2
    foot.add_run('10 September 2026  •  ')
    fld=OxmlElement('w:fldSimple'); fld.set(qn('w:instr'),'PAGE'); foot._p.append(fld)
    lines=md.splitlines(); i=0; in_refs=False
    while i<len(lines):
        line=lines[i].strip(); i+=1
        if not line: continue
        if line.startswith('# '): doc.add_paragraph(line[2:],'Title'); continue
        if line.startswith('## '):
            title=line[3:]
            if title in ['Abstract','Tables','Figure legends','References']: doc.add_page_break()
            in_refs=title=='References'; doc.add_heading(title,1); continue
        if line.startswith('### '):
            if line.startswith('### Figure 2.'): doc.add_page_break()
            doc.add_heading(line[4:],2); continue
        if line.startswith('![Figure'):
            name=re.search(r'\(([^)]+)\)',line).group(1)
            doc.add_picture(str(ROOT/name),width=Inches(6.45)); continue
        if line.startswith('|'):
            rows=[line]
            while i<len(lines) and lines[i].strip().startswith('|'): rows.append(lines[i].strip()); i+=1
            content=[[c.strip() for c in r.strip('|').split('|')] for r in rows if not re.fullmatch(r'[| :\-]+',r)]
            tab=doc.add_table(rows=1,cols=len(content[0])); tab.style='Light Shading Accent 1'; tab.autofit=False
            # Per-table column widths, keyed by column count, with an even
            # split as the fallback. The two hardcoded layouts covered Table 1
            # and Table 2 only, so adding a five-column comparison table raised
            # an IndexError from the width lookup rather than laying the table
            # out badly -- which is the better failure, but still a failure.
            n=len(content[0])
            widths={7:[1.55,.6,1.0,.5,1.,1.,.8],
                    5:[1.35,2.15,1.15,1.05,1.25],
                    3:[1.2,2.25,3.0]}.get(n,[6.45/n]*n)
            for col,width in zip(tab.columns,widths): col.width=Inches(width)
            for j,c in enumerate(content[0]): tab.rows[0].cells[j].text=c
            header=OxmlElement('w:tblHeader'); tab.rows[0]._tr.get_or_add_trPr().append(header)
            for vals in content[1:]:
                cells=tab.add_row().cells
                for j,c in enumerate(vals): cells[j].text=c
            for row in tab.rows:
                no_split=OxmlElement('w:cantSplit'); row._tr.get_or_add_trPr().append(no_split)
                for column,cell in enumerate(row.cells):
                    cell.width=Inches(widths[column])
                    for p in cell.paragraphs:
                        p.paragraph_format.space_after=Pt(4); p.paragraph_format.space_before=Pt(3)
                        p.paragraph_format.line_spacing=1.0
                        for r in p.runs: r.font.name='Arial'; r.font.size=Pt(9)
            doc.add_paragraph(); continue
        if in_refs and line.startswith('10. '): doc.add_page_break()
        p=doc.add_paragraph(); add_text(p,line)
        if in_refs:
            p.paragraph_format.line_spacing=1.0; p.paragraph_format.space_after=Pt(8)
            p.paragraph_format.keep_together=True
            for r in p.runs: r.font.size=Pt(9.5)
    doc.core_properties.title=md.splitlines()[0][2:]
    doc.core_properties.subject='NAR Web Server Issue manuscript draft with retained internal validation'
    doc.core_properties.author=('Auroni Deep; Himanshu Shekhar; Shilpi Minocha; '
                                'Vivekanandan Perumal; Saran Kumar')
    doc.core_properties.keywords='QuadCond, AENNA-3D, G-quadruplex, i-motif'
    target = ROOT / 'QuadCond_AENNA_NAR_manuscript.docx'
    # Word holds an exclusive lock and leaves a `~$` sidecar while a document is
    # open. Overwriting under it either fails or, worse, races the copy the
    # author is editing. The build writes a clearly-named sibling instead and
    # says so, rather than dying at the last line after doing all the work.
    lock = target.with_name('~$' + target.name[2:])
    try:
        if lock.exists():
            raise PermissionError('open in Word')
        doc.save(target)
        saved = target
    except PermissionError:
        saved = target.with_name(target.stem + '.rebuilt.docx')
        doc.save(saved)
        print(f'!! {target.name} is open in Word; wrote {saved.name} instead. '
              f'Close the document and re-run, or rename the rebuilt file.')
    return saved


def main():
    workflow(); ablation()
    regression=[]
    for key,label in zip(HEADS,LABELS):
        h=MAIN['heads'][key]; t=h['training_meta']; m=h['metrics']
        unit='°C' if '_tm' in key else 'pH units'
        regression.append([label,f"{t['n_rows']:,}",f"{t['grouping']['n_groups']} {t['group_by']}",
          f"{m['r2']:.3f}",f"{m['rmse']:.3f} {unit}",f"{m['mae']:.3f} {unit}",f"{m['split_conformal_coverage']:.3f}"])
    hdr=['Target','n','Groups','R²','RMSE','MAE','Coverage']
    csvout('Table_1_regression.csv',hdr,regression)
    supp=[]
    for key,h in MAIN['heads'].items():
        t=h['training_meta']; m=h['metrics']
        supp.append([key,t['n_rows'],t['group_by'],t['grouping']['n_groups'],t['n_folds'],t['n_seeds'],
          t.get('target_semantics',''),json.dumps(t.get('label_classes',{})),json.dumps(t.get('sources',{})),json.dumps(m)])
    csvout('Table_S1_all_heads.csv',['head','records','group_by','groups','folds','seeds','target_semantics','label_classes','sources','metrics_json'],supp)
    summary=MAIN['atlas_snapshot']['summary']; keys=list(summary[0])
    csvout('Table_S2_atlas_sources.csv',keys,[[s.get(k,'') for k in keys] for s in summary])
    text=(ROOT/'manuscript_source.md').read_text(encoding='utf-8')
    text=text.replace('{{REGRESSION_TABLE}}',table_md(hdr,regression))
    text=text.replace('{{BENCHMARK_TABLE}}',benchmark_table())
    text=text.replace('{{SCANNER_TABLE}}',scanner_table())
    text=text.replace('{{GENOMEWIDE_TABLE}}',genomewide_table())
    for n,stem in [(1,'Figure_1_workflow'),(2,'Figure_2_ablation'),(3,'Figure_3_g4_folding'),
                   (4,'Figure_4_topology_tm_im'),(5,'Figure_5_variant_dtm'),(6,'Figure_6_condition_pht')]:
        text=text.replace('{{FIGURE_%d}}'%n,'![Figure %d](figures/%s.png)'%(n,stem))
    text=references(text)
    assert not re.search(r'\{\{|\[@',text)
    (ROOT/'QuadCond_AENNA_NAR_manuscript.md').write_text(text,encoding='utf-8')
    docx_from_md(text)
    provenance={
        'software_revision':'0.5.1','frozen_model_version':MAIN['version'],
        'model_artifact_sha256':MAIN['artifact_sha256'],'dataset_fingerprint':MAIN['dataset_fingerprint_sha256'],
        'inputs':{p.name:hashlib.sha256(p.read_bytes()).hexdigest() for p in [ART/'quadcond_model.json',ART/'quadcond_model_seqonly.json']},
        'method':'Copy retained metadata and plot it; no model fitting or performance recomputation.',
        'body_word_count_before_tables':len(text.split('## Tables')[0].split()),
        'external_validation_claimed':False,'public_deployment_verified':False}
    (ROOT/'manuscript_provenance.json').write_text(json.dumps(provenance,indent=2)+'\n',encoding='utf-8')
    audit=json.loads((ROOT/'reference_audit.json').read_text(encoding='utf-8'))
    # Counted, not typed. The literal 18 stayed correct only as long as nobody
    # added a reference, which is the shortest-lived kind of true statement in
    # a build script.
    print(f"Created manuscript; {provenance['body_word_count_before_tables']} words before tables; "
          f"{audit['count']} verified references; "
          f"{len(audit['locators_to_complete_before_submission'])} locators still to complete.")


if __name__=='__main__': main()
